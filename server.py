# -*- coding: utf-8 -*-
"""
A股人气雷达 - 后端代理服务器
"""

import io
import os
import sys
import json
import time
import hashlib
import requests
import urllib3
from datetime import datetime, timedelta
from flask import Flask, jsonify, send_from_directory, send_file, request
from flask_cors import CORS
from stock_data import get_stock_name, STOCK_NAMES
import theme_heat_service as _theme_heat
import market_db as _mdb

# 禁用SSL警告
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

app = Flask(__name__, static_folder='.', static_url_path='')
CORS(app)

# 内存缓存
cache = {}
CACHE_TTL = 600  # 10分钟

# ==================== 交易日历 ====================

_CAL_CACHE = {"days": [], "ts": 0.0, "ttl": 1800}
_CAL_TTL = 1800  # 交易日历缓存 30 分钟


def _trading_calendar(n=60):
    """真实交易日历（倒序，最新在前，天然跳过周末与节假日）。

    数据来源：上证指数日K（screener.get_trading_days）。
    取数失败时返回空列表，由调用方走兜底逻辑。
    """
    now_ts = time.time()
    cached = _CAL_CACHE["days"]
    if cached and (now_ts - _CAL_CACHE["ts"]) < _CAL_CACHE.get("ttl", _CAL_TTL):
        return cached[:n]

    want = max(n, 60)
    today = datetime.now().strftime('%Y-%m-%d')
    is_weekday = datetime.now().isoweekday() <= 5
    # 1) 读库优先
    days = _mdb.load_calendar(want)
    need_net = True
    if days:
        try:
            gap = (datetime.now().date() - datetime.strptime(days[0], '%Y-%m-%d').date()).days
        except Exception:
            gap = 999
        # 库中最新交易日够新 且 条数够 -> 直接用库
        # 关键：工作日若库里最新交易日 < 今天，必须联网刷新。
        # 否则「新交易日的第一天」会被当成非交易日：场次日期停在昨天，
        # 且定时任务（tick 里 today not in days 就 return）会被整体跳过。
        stale_today = (gap >= 1) and is_weekday
        need_net = (gap > 5) or (len(days) < want) or stale_today
    # 2) 需要时走接口，并落库
    if need_net:
        try:
            import screener as _sc
            api_days = [d for d in _sc.recent_trading_days(want) if d]
        except Exception:
            api_days = []
        if api_days:
            _mdb.save_calendar(api_days)
            days = api_days
    if days:
        _CAL_CACHE["days"] = days
        _CAL_CACHE["ts"] = now_ts
        # 工作日若日历里仍没有今天（数据源今天还没出 / 今天休市）-> 短缓存，稍后自动重试
        _CAL_CACHE["ttl"] = 300 if (is_weekday and today not in set(days)) else _CAL_TTL
        return days[:n]
    return cached[:n] if cached else []


def _fallback_last_trading_day(date):
    """兜底：仅按周末回溯（交易日历不可用时）"""
    for i in range(30):
        check_date = date - timedelta(days=i)
        if check_date.weekday() < 5:
            return check_date.strftime('%Y-%m-%d')
    return (date - timedelta(days=1)).strftime('%Y-%m-%d')


def _last_completed_trading_day(now=None):
    """最近一个【已收盘】的交易日（含节假日回溯）。

    - 交易日 15:00 后        -> 当日
    - 交易日盘中/开盘前       -> 上一交易日
    - 周末/节假日            -> 最近的历史交易日
    """
    now = now or datetime.now()
    today = now.strftime('%Y-%m-%d')
    days = _trading_calendar(60)            # 倒序：最新在前
    if not days:
        return _fallback_last_trading_day(now)
    closed_today = (now.hour * 60 + now.minute) >= 900
    for d in days:
        if d < today:
            return d
        if d == today and closed_today:
            return d
    return days[-1]


def get_last_trading_day(date=None):
    """获取 date 当天或之前最近的已完成交易日（含节假日回溯）"""
    if date is None:
        date = datetime.now()
    days = _trading_calendar(60)
    if not days:
        return _fallback_last_trading_day(date)
    target = date.strftime('%Y-%m-%d')
    for d in days:                          # 倒序
        if d <= target:
            return d
    return days[-1]


def get_data_trading_day():
    """根据当前时刻确定数据所属交易日"""
    now = datetime.now()
    today = now.strftime('%Y-%m-%d')
    days = _trading_calendar(60)
    last = _last_completed_trading_day(now)

    if last is None:
        return today, ""
    if last == today:
        return today, ""

    # 生成说明文案
    if today in set(days):
        if now.hour * 60 + now.minute < 570:
            return last, "今日未开盘，以下为上一交易日收盘数据"
        return last, "今日盘中，以下为上一交易日收盘数据"
    return last, "今日为节假日/非交易日，已回溯至最近交易日"



def _market_is_live():
    """当前是否处于当日交易时段（此时需要实时数据，不走库）"""
    now = datetime.now()
    today = now.strftime('%Y-%m-%d')
    if today not in set(_trading_calendar(60)):
        return False
    m = now.hour * 60 + now.minute
    return (9 * 60 + 15) <= m <= (15 * 60 + 5)


def _session_date():
    """当前数据所属交易日：交易日=今天；否则=最近已完成交易日"""
    now = datetime.now()
    today = now.strftime('%Y-%m-%d')
    if today in set(_trading_calendar(60)):
        return today
    return _last_completed_trading_day(now)


def attach_streaks(rows, trade_date, top_n=30, code_key="code"):
    """给榜单每行补上「连续上榜天数」（含当天，按交易日历往前回推）

    code_key: 取哪个字段作为股票代码（人气榜用 code，题材个股用 ticker）
    """
    if not rows or not trade_date:
        return rows or []
    try:
        streaks = _mdb.hotlist_streaks(trade_date, top_n=top_n)
    except Exception:
        streaks = {}
    for r in rows:
        code = (r.get(code_key) or "").replace("SH", "").replace("SZ", "").replace("BJ", "")
        r["streak_days"] = streaks.get(code, 1)
    return rows


def build_hotlist_live():
    """实时抓取人气榜。返回 (rows, source, degraded, error)。

    供 /api/hotlist 与定时任务共用；error 非空表示两个数据源都失败。
    """
    stock_list, error = fetch_ths_hotlist()
    if error or not stock_list:
        # 降级到东方财富
        eastmoney_data, error2 = fetch_eastmoney_hotlist()
        if not eastmoney_data:
            return None, None, False, (error or error2 or "未知错误")
        codes = [item.get('sc', '') for item in eastmoney_data[:30]]
        quote_data, _ = fetch_eastmoney_batch_quote(codes)
        return process_eastmoney_data(eastmoney_data, quote_data), "东方财富热股榜", True, None
    codes = [item.get('code', '') for item in stock_list[:30]]
    quote_data, _ = fetch_eastmoney_batch_quote(codes)
    return process_ths_data(stock_list, quote_data), "同花顺人气榜", False, None


def _hotlist_summary(result_list):
    """人气榜汇总（实时构建与库回放共用）"""
    total_count = len(result_list)
    if total_count == 0:
        return {'main_direction': '暂无明显主线方向', 'hot_proportion': '无数据'}
    hot_count = sum(1 for item in result_list if item.get('is_hot'))
    if hot_count > total_count * 0.6:
        summary_hot = "涨停股占据多数"
    elif hot_count > total_count * 0.3:
        summary_hot = "涨停股与非涨停股参半"
    else:
        summary_hot = "多数为非涨停股"
    all_concepts = []
    for item in result_list:
        c = item.get('concept_tag')
        if c and c not in ('未获取', '降级模式，无数据', '—', '-', ''):
            all_concepts.extend([x.strip() for x in c.split(',')])
    if all_concepts:
        from collections import Counter
        top_concept = Counter(all_concepts).most_common(1)[0][0]
        summary_main = "人气主线集中在%s方向" % top_concept
    else:
        summary_main = "暂无明显主线方向"
    return {'main_direction': summary_main, 'hot_proportion': summary_hot}


# ==================== 缓存 ====================

def get_cache_key(url):
    return hashlib.md5(url.encode()).hexdigest()

def get_cached_data(key):
    if key in cache:
        data, timestamp = cache[key]
        if time.time() - timestamp < CACHE_TTL:
            return data
        else:
            del cache[key]
    return None

def set_cached_data(key, data):
    cache[key] = (data, time.time())


# ==================== 数据获取 ====================

def fetch_ths_hotlist():
    """获取同花顺人气榜数据"""
    url = "https://dq.10jqka.com.cn/fuyao/hot_list_data/out/hot_list/v1/stock?stock_type=a&type=hour&list_type=normal"
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Referer': 'https://www.10jqka.com.cn/',
        'Accept': 'application/json, text/plain, */*',
    }
    
    try:
        response = requests.get(url, headers=headers, timeout=15, verify=False)
        response.raise_for_status()
        data = response.json()
        
        # 同花顺返回结构: {"status_code": 0, "data": {"stock_list": [...]}, "status_msg": "success"}
        stock_list = data.get('data', {}).get('stock_list', [])
        
        if stock_list and len(stock_list) > 0:
            return stock_list[:30], None
        else:
            return None, f"同花顺接口返回数据异常: {data.get('status_msg', '未知错误')}"
    except Exception as e:
        return None, f"同花顺接口请求失败: {str(e)}"


def fetch_eastmoney_hotlist():
    """获取东方财富热股榜"""
    url = "https://emappdata.eastmoney.com/stockrank/getAllCurrentList"
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
        'Content-Type': 'application/json',
    }
    payload = {
        "appId": "appId01",
        "globalId": "786e4c21-70dc-435a-93bb-38",
        "marketType": "",
        "pageNo": 1,
        "pageSize": 30
    }
    
    try:
        response = requests.post(url, headers=headers, json=payload, timeout=15, verify=False)
        response.raise_for_status()
        data = response.json()
        
        if 'data' in data and len(data['data']) > 0:
            return data['data'], None
        else:
            return None, "东方财富接口返回数据异常"
    except Exception as e:
        return None, f"东方财富接口请求失败: {str(e)}"


def fetch_eastmoney_batch_quote(codes):
    """批量获取东方财富实时行情（含名称和行业）"""
    secids = []
    for code in codes:
        clean = code.replace('SH','').replace('SZ','').replace('BJ','')
        if clean.startswith('6'):
            secids.append(f"1.{clean}")
        else:
            secids.append(f"0.{clean}")
    
    secids_str = ','.join(secids)
    url = f"https://push2.eastmoney.com/api/qt/ulist.np/get?secids={secids_str}&fields=f12,f14,f100&fltt=2"
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)',
        'Referer': 'https://www.eastmoney.com/'
    }
    
    try:
        response = requests.get(url, headers=headers, timeout=10, verify=False)
        response.raise_for_status()
        data = response.json()
        
        result = {}
        if data and data.get('data') and data['data'].get('diff'):
            for item in data['data']['diff']:
                code = item.get('f12', '')
                name = item.get('f14', '')
                industry = item.get('f100', '')
                if code:
                    result[code] = {'name': name, 'industry': industry}
        return result, None
    except Exception as e:
        return {}, str(e)


# ==================== 数据处理 ====================

def get_secid_prefix(code):
    clean_code = code.replace('SH', '').replace('SZ', '').replace('BJ', '')
    if clean_code.startswith('6'):
        return f"1.{clean_code}"
    else:
        return f"0.{clean_code}"


def process_consecutive_boards(popularity_tag):
    if not popularity_tag:
        return "无"
    
    import re
    match = re.search(r'(\d+)天(\d+)板', popularity_tag)
    if match:
        days = match.group(1)
        boards = match.group(2)
        return f"{boards}连板（{days}天{boards}板）"
    
    if '首板' in popularity_tag:
        return "首板（当日）"
    
    return "无"


def get_tier(order):
    if order <= 10:
        return "第一梯队·人气Top10"
    elif order <= 20:
        return "第二梯队·人气Top20"
    else:
        return "第三梯队·人气Top30"


def get_anomaly_analysis(item):
    if item.get('analyse_title'):
        return item['analyse_title'][:30]
    
    if item.get('analyse'):
        text = item['analyse'].replace('\n', ' ').strip()
        return text[:30]
    
    if item.get('topic') and item['topic'].get('title'):
        return item['topic']['title'][:30]
    
    return "无"


def process_ths_data(stock_list, quote_data):
    result = []
    
    for item in stock_list:
        code = item.get('code', '')
        clean_code = code.replace('SH', '').replace('SZ', '').replace('BJ', '')
        name = item.get('name', '')
        order = item.get('order', 0)
        popularity_tag = item.get('tag', {}).get('popularity_tag', '') if item.get('tag') else ''
        
        # 从行情数据补充名称和行业
        quote = quote_data.get(clean_code, {})
        if not name and quote.get('name'):
            name = quote['name']
        elif not name:
            name = get_stock_name(clean_code)
        
        # 从 concept_tag 提取概念标签
        concept_tags = item.get('tag', {}).get('concept_tag', []) if item.get('tag') else []
        if concept_tags and isinstance(concept_tags, list):
            concept_str = ', '.join(concept_tags[:3])  # 最多取3个
        else:
            concept_str = '未获取'
        
        # 获取涨跌幅
        rise_and_fall = item.get('rise_and_fall', 0)
        
        result.append({
            'rank': order,
            'name': name or clean_code,
            'code': code,
            'rise_and_fall': rise_and_fall,
            'consecutive_boards': process_consecutive_boards(popularity_tag),
            'tier': get_tier(order),
            'concept_tag': concept_str,
            'anomaly_analysis': get_anomaly_analysis(item),
            'is_hot': bool(popularity_tag)
        })
    
    return result


def process_eastmoney_data(data_list, quote_data):
    result = []
    
    for i, item in enumerate(data_list[:30], 1):
        code = item.get('sc', '')
        clean_code = code.replace('SH', '').replace('SZ', '').replace('BJ', '')
        
        # 从行情数据获取名称
        quote = quote_data.get(clean_code, {})
        name = quote.get('name', '') or get_stock_name(clean_code) or clean_code
        
        result.append({
            'rank': i,
            'name': name,
            'code': code,
            'rise_and_fall': 0,
            'consecutive_boards': "无",
            'tier': get_tier(i),
            'concept_tag': "降级模式，无数据",
            'anomaly_analysis': "无",
            'is_hot': False
        })
    
    return result


# ==================== API路由 ====================

@app.route('/')
def index():
    return send_from_directory('.', 'index.html')


@app.route('/api/hotlist')
def get_hotlist():
    trading_day, message = get_data_trading_day()
    user_date = (request.args.get('date') or '').strip()

    # ---- ① 明确指定了日期：只读本地库；读不到就如实告知，绝不回退成实时数据 ----
    if user_date:
        db_hot = _mdb.load_hotlist(user_date)
        if db_hot and db_hot.get('data'):
            return jsonify({
                'success': True,
                'trading_day': user_date,
                'message': '' if user_date == trading_day else '历史数据（来自本地数据库）',
                'source': db_hot.get('source') or '本地数据库',
                'degraded': bool(db_hot.get('degraded')),
                'collection_time': db_hot.get('captured_at'),
                'data': attach_streaks(db_hot['data'], user_date),
                'summary': _hotlist_summary(db_hot['data']),
                'from_db': True,
            })
        return jsonify({
            'success': True, 'empty': True,
            'trading_day': user_date,
            'message': '该交易日没有已保存的人气榜数据（当天程序可能没有运行）',
            'source': '', 'degraded': False, 'collection_time': None,
            'data': [], 'summary': {'main_direction': '', 'hot_proportion': ''},
            'from_db': True,
        })

    # ---- ② 没指定日期：优先用本地库里「最近已完成交易日」的快照 ----
    # 重要：**盘中也要优先读库**。因为数据源只给"此刻"的实时榜，若盘中抓取
    # 并归到上一交易日，就会把上一交易日的收盘快照污染掉（改不回来）。
    # 本应用的定位本来就是"收盘后数据"，盘中显示上一交易日收盘快照是正确的。
    if True:
        db_hot = _mdb.load_hotlist(trading_day)
        if db_hot and db_hot.get('data'):
            return jsonify({
                'success': True,
                'trading_day': trading_day,
                'message': message,
                'source': db_hot.get('source') or '本地数据库',
                'degraded': bool(db_hot.get('degraded')),
                'collection_time': db_hot.get('captured_at'),
                'data': attach_streaks(db_hot['data'], trading_day),
                'summary': _hotlist_summary(db_hot['data']),
                'from_db': True,
            })

    cache_key = get_cache_key('hotlist')
    cached = get_cached_data(cache_key)
    if cached:
        return jsonify(cached)

    # 实时抓取（同花顺优先，失败降级东方财富）
    result_list, source, degraded, _err = build_hotlist_live()
    if result_list is None:
        return jsonify({
            'success': False,
            'error': f"数据获取失败: {_err}"
        }), 500
    
    # 生成总结
    _summary = _hotlist_summary(result_list)
    summary_main = _summary['main_direction']
    summary_hot = _summary['hot_proportion']
    collection_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    # 实时抓到的其实是"当前场次"的榜，绝不能记到上一交易日名下（会污染历史）
    _save_date = _session_date()
    attach_streaks(result_list, _save_date)

    response_data = {
        'success': True,
        'trading_day': _save_date,
        'message': (message if _save_date == trading_day
                    else '以下为实时数据（本地库暂无该交易日快照）'),
        'source': source,
        'degraded': degraded,
        'collection_time': collection_time,
        'data': result_list,
        'summary': {
            'main_direction': summary_main,
            'hot_proportion': summary_hot
        }
    }
    
    set_cached_data(cache_key, response_data)
    _mdb.save_hotlist(_save_date, result_list, source=source, degraded=degraded,
                      captured_at=collection_time)
    _mdb.log_fetch('hotlist', source=source, trade_date=_save_date, rows=len(result_list),
                   status='degraded' if degraded else 'ok')
    return jsonify(response_data)



# ==================== 图表数据API ====================

def get_secid(code):
    """将股票代码转换为东方财富secid格式"""
    clean = code.replace('SH', '').replace('SZ', '').replace('BJ', '')
    if clean.startswith('6'):
        return f"1.{clean}"
    else:
        return f"0.{clean}"




@app.route('/api/stock/<code>/news')
def get_stock_news(code):
    clean_code = code.replace('SH', '').replace('SZ', '').replace('BJ', '')
    try:
        url = 'https://np-anotice-stock.eastmoney.com/api/security/ann?sr=-1&page_size=8&page_index=1&ann_type=A&client_source=web&f_node=0&s_node=0&stock_list=' + clean_code
        headers = {'User-Agent': 'Mozilla/5.0', 'Referer': 'https://data.eastmoney.com/'}
        response = requests.get(url, headers=headers, timeout=10, verify=False)
        data = response.json()
        
        result = []
        if data.get('data') and data['data'].get('list'):
            for item in data['data']['list'][:8]:
                result.append({
                    'title': item.get('title', ''),
                    'date': item.get('notice_date', '')[:10] if item.get('notice_date') else '',
                    'source': '东方财富公告',
                    'url': 'https://data.eastmoney.com/notices/detail/' + clean_code + '/' + item.get('art_code', '') + '.html'
                })
        
        if result:
            _mdb.save_news(clean_code, result)
            _mdb.log_fetch('news', source='eastmoney', rows=len(result), status='ok')
            return jsonify({'success': True, 'data': result})
        db_news = _mdb.load_news(clean_code, 8)
        if db_news:
            return jsonify({'success': True, 'data': db_news, 'from_db': True})
        return jsonify({'success': True, 'data': result})
    except Exception as e:
        db_news = _mdb.load_news(clean_code, 8)
        if db_news:
            return jsonify({'success': True, 'data': db_news, 'from_db': True})
        return jsonify({'success': False, 'error': str(e)})

@app.route('/api/stock/<code>')
def get_stock_chart(code):
    """获取个股分时或日K数据"""
    chart_type = request.args.get('type', 'minute')
    clean_code = code.replace('SH', '').replace('SZ', '').replace('BJ', '')
    
    # 判断市场前缀
    if clean_code.startswith('6'):
        market_code = f"sh{clean_code}"
    else:
        market_code = f"sz{clean_code}"
    
    try:
        if chart_type == 'minute':
            # ---- 非交易时段：读库优先 ----
            if not _market_is_live():
                _sess_m = _session_date()
                _db_min = _mdb.load_minute(clean_code, _sess_m)
                if len(_db_min) >= 30:
                    _pc = _mdb.load_prev_close(clean_code, _sess_m) or 0
                    return jsonify({'success': True, 'data': _db_min,
                                    'prev_close': _pc, 'from_db': True})
            # 使用腾讯接口获取分时数据
            url = f"https://web.ifzq.gtimg.cn/appstock/app/minute/query?code={market_code}"
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)',
                'Referer': 'https://finance.qq.com/'
            }
            response = requests.get(url, headers=headers, timeout=15, verify=False)
            data = response.json()
            
            # 腾讯返回结构: data -> sh600664 -> data -> data (数组)
            stock_data = data.get('data', {}).get(market_code, {})
            minute_data = stock_data.get('data', {}).get('data', [])
            
            if minute_data:
                result = []
                for item in minute_data:
                    parts = item.split(' ')
                    if len(parts) >= 3:
                        time_str = parts[0]
                        # 格式化时间 HHMM -> HH:MM
                        if len(time_str) == 4:
                            time_str = time_str[:2] + ':' + time_str[2:]
                        result.append({
                            'time': time_str,
                            'price': float(parts[1]) if parts[1] else 0,
                            'volume': int(parts[2]) if parts[2] else 0
                        })
                # 获取前一日收盘价（使用日K接口）
                prev_close = 0
                try:
                    kline_url = f"https://quotes.sina.cn/cn/api/jsonp_v2.php/var/CN_MarketDataService.getKLineData?symbol={market_code}&scale=240&ma=no&datalen=2"
                    kline_resp = requests.get(kline_url, headers=headers, timeout=10, verify=False)
                    kline_text = kline_resp.text
                    kline_start = kline_text.find('(')
                    kline_end = kline_text.rfind(')')
                    if kline_start >= 0 and kline_end > kline_start:
                        import json as json_lib
                        klines = json_lib.loads(kline_text[kline_start+1:kline_end])
                        if len(klines) >= 2:
                            prev_close = float(klines[-2].get('close', 0))
                except:
                    pass
                
                _sess_log = _session_date()
                _mdb.save_minute(clean_code, _sess_log, result)
                _mdb.log_fetch('minute_kline', source='tencent', trade_date=_sess_log,
                               rows=len(result), status='ok')
                return jsonify({'success': True, 'data': result, 'prev_close': prev_close})
            else:
                return jsonify({'success': False, 'error': '无分时数据'})
                
        else:
            # ---- 非交易时段：读库优先 ----
            if not _market_is_live():
                _db_k = _mdb.load_daily_kline(clean_code, limit=60)
                if len(_db_k) >= 60:
                    return jsonify({'success': True, 'data': _db_k, 'from_db': True})
            # 使用新浪K线接口（更稳定）
            url = f"https://quotes.sina.cn/cn/api/jsonp_v2.php/var/CN_MarketDataService.getKLineData?symbol={market_code}&scale=240&ma=no&datalen=60"
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)',
                'Referer': 'https://finance.sina.com.cn/'
            }
            response = requests.get(url, headers=headers, timeout=15, verify=False)
            text = response.text
            
            # 解析JSONP响应
            start = text.find('(')
            end = text.rfind(')')
            
            if start >= 0 and end > start:
                import json as json_lib
                kline_data = json_lib.loads(text[start+1:end])
                
                if kline_data:
                    result = []
                    for item in kline_data:
                        result.append({
                            'date': item.get('day', ''),
                            'open': float(item.get('open', 0)),
                            'close': float(item.get('close', 0)),
                            'high': float(item.get('high', 0)),
                            'low': float(item.get('low', 0)),
                            'volume': int(item.get('volume', 0))
                        })
                    _mdb.save_daily_kline(clean_code, result)
                    _mdb.log_fetch('daily_kline', source='sina', trade_date=result[-1].get('date'),
                                   rows=len(result), status='ok')
                    return jsonify({'success': True, 'data': result})
            
            return jsonify({'success': False, 'error': '无日K数据'})
                
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


# ==================== 动态选股 API ====================

import threading as _threading
import uuid as _uuid
import screener as _screener

SCREEN_JOBS = {}
SCREEN_CACHE = {}


def _screen_worker(job_id, target):
    job = SCREEN_JOBS.get(job_id)
    if job is None:
        return

    def cb(done, total, msg):
        job["done"] = done
        job["total"] = total
        job["message"] = msg

    try:
        res = _screener.run_screen(target, progress=cb)
        job["result"] = res
        job["state"] = "done"
        if res.get("success"):
            SCREEN_CACHE[target] = res
            _mdb.save_screening(res)
    except Exception as e:
        job["result"] = {"success": False, "error": "扫描异常: %s" % e}
        job["state"] = "error"
    finally:
        job["finishedAt"] = time.time()


@app.route("/api/trading-days")
def trading_days_api():
    try:
        n = int(request.args.get("n", 10))
    except Exception:
        n = 10
    n = max(1, min(n, 60))
    days = _trading_calendar(n)          # 内部：读库优先 + 新鲜度校验 + 落库
    if not days:
        days = _screener.recent_trading_days(n)
    return jsonify({"success": True, "days": days})


@app.route("/api/screen/start")
def screen_start():
    target = (request.args.get("date") or "").strip()
    if not target:
        return jsonify({"success": False, "error": "missing param: date"})
    if target in SCREEN_CACHE:
        return jsonify({"success": True, "cached": True, "result": SCREEN_CACHE[target]})
    db_res = _mdb.load_screening(target)
    if db_res:
        db_res["from_db"] = True
        SCREEN_CACHE[target] = db_res
        return jsonify({"success": True, "cached": True, "result": db_res, "from_db": True})
    for jid, j in SCREEN_JOBS.items():
        if j.get("target") == target and j.get("state") == "running":
            return jsonify({"success": True, "jobId": jid, "cached": False})
    jid = _uuid.uuid4().hex[:12]
    SCREEN_JOBS[jid] = {
        "target": target,
        "state": "running",
        "done": 0,
        "total": 0,
        "message": "preparing...",
        "result": None,
        "startedAt": time.time(),
    }
    _threading.Thread(target=_screen_worker, args=(jid, target), daemon=True).start()
    return jsonify({"success": True, "jobId": jid, "cached": False})


@app.route("/api/screen/status")
def screen_status():
    jid = request.args.get("jobId", "")
    job = SCREEN_JOBS.get(jid)
    if not job:
        return jsonify({"success": False, "error": "job not found"})
    out = {
        "success": True,
        "state": job["state"],
        "done": job["done"],
        "total": job["total"],
        "message": job["message"],
    }
    if job["state"] in ("done", "error"):
        out["result"] = job["result"]
    return jsonify(out)


@app.route("/api/screen/report")
def screen_report():
    target = (request.args.get("date") or "").strip()
    res = SCREEN_CACHE.get(target) or _mdb.load_screening(target)
    if not res:
        return "该日期还没有扫描结果，请先在页面完成一次扫描。", 404
    return _screener.render_report_html(res), 200, {"Content-Type": "text/html; charset=utf-8"}


@app.route("/api/screen")
def screen_sync():
    target = (request.args.get("date") or "").strip()
    if not target:
        return jsonify({"success": False, "error": "missing param: date"})
    if target in SCREEN_CACHE:
        return jsonify(SCREEN_CACHE[target])
    db_res = _mdb.load_screening(target)
    if db_res:
        db_res["from_db"] = True
        SCREEN_CACHE[target] = db_res
        return jsonify(db_res)
    res = _screener.run_screen(target)
    if res.get("success"):
        SCREEN_CACHE[target] = res
        _mdb.save_screening(res)
    return jsonify(res)


def _fill_theme_change(res, trade_date):
    """给题材 Top50 补「涨跌幅」：本地库回放时用日K，实时构建时数据已自带"""
    rows = (res or {}).get("topStocks") or []
    if not rows or not trade_date:
        return res
    if all(r.get("change") is not None for r in rows):
        return res
    try:
        cmap = _mdb.kline_change_map(trade_date)
    except Exception:
        cmap = {}
    if cmap:
        for r in rows:
            if r.get("change") is None:
                v = cmap.get(str(r.get("ticker") or ""))
                if v is not None:
                    r["change"] = v
    return res


@app.route("/api/event-themes")
def event_themes_api():
    """当日事件驱动题材 + 人气排序"""
    import event_theme_service as _event_theme
    try:
        limit = int(request.args.get("limit", 6))
    except Exception:
        limit = 6
    try:
        stocks = int(request.args.get("stocks", 10))
    except Exception:
        stocks = 10
    # ---- ① 明确指定日期：只从本地库回放 ----
    req_date = (request.args.get("date") or "").strip()
    if req_date:
        db_res = _mdb.load_event_themes(req_date)
        if db_res and db_res.get("success"):
            db_res["from_db"] = True
            db_res["trade_date"] = req_date
            _fill_theme_change(db_res, req_date)
            attach_streaks(db_res.get("topStocks") or [], req_date, code_key="ticker")
            return jsonify(db_res)
        # 用 200 + 明确错误码，避免浏览器控制台出现无意义的 404
        return jsonify({"success": False, "error": "NO_DATA",
                        "trade_date": req_date,
                        "message": "该交易日没有已保存的题材数据（当天程序可能没有运行）"})

    live = _market_is_live()
    _sess = _session_date()

    # ---- ② 非交易时段/非交易日：读库优先（不受 limit/stocks 影响），且绝不用实时重建覆盖存档 ----
    #     重要：数据源只给"此刻"的题材/事件；若在非交易时段重建并归到「上一交易日」，
    #     会把该交易日的收盘存档污染掉（不可恢复）。所以非交易时段**只读不写**。
    if not live:
        _db_theme = _mdb.load_event_themes(_sess)
        if _db_theme and _db_theme.get("success"):
            _db_theme["from_db"] = True
            _db_theme["trade_date"] = _sess
            _fill_theme_change(_db_theme, _sess)
            attach_streaks(_db_theme.get("topStocks") or [], _sess, code_key="ticker")
            return jsonify(_db_theme)
    try:
        res = _event_theme.build_event_themes(limit, stocks)
        if res.get("success"):
            # 只有「场次日期就是今天」才落库：
            #   · 交易日（盘中 / 盘后）-> _sess == 今天 -> 正常存档
            #   · 周末 / 节假日        -> _sess == 上一交易日 -> 不写，避免覆盖历史存档
            if _sess == datetime.now().strftime("%Y-%m-%d"):
                _mdb.save_event_themes(res, trade_date=_sess)
            attach_streaks(res.get("topStocks") or [], _sess, code_key="ticker")
        return jsonify(res)
    except Exception as e:
        return jsonify({"success": False, "error": "event themes failed: %s" % e}), 500


@app.route("/api/theme-hot")
def theme_hot_api():
    """自动发现近期热门题材"""
    try:
        limit = int(request.args.get("limit", 12))
    except Exception:
        limit = 12
    try:
        lookback = int(request.args.get("lookback", 5))
    except Exception:
        lookback = 5
    try:
        return jsonify(_theme_heat.discover_hot_themes(limit, lookback))
    except Exception as e:
        return jsonify({"success": False, "error": "theme discovery failed: %s" % e}), 500


@app.route("/api/theme-heat")
def theme_heat_api():
    theme = (request.args.get("theme") or "").strip()
    try:
        limit = int(request.args.get("limit", 60))
    except Exception:
        limit = 60
    if not theme:
        return jsonify({"success": False, "error": "missing param: theme"})
    try:
        return jsonify(_theme_heat.query_theme_heat(theme, limit))
    except Exception as e:
        return jsonify({"success": False, "error": "theme heat query failed: %s" % e}), 500

@app.route("/api/db/overview")
def db_overview_api():
    """数据管理页：各表行数与日期范围"""
    try:
        return jsonify({"success": True, "overview": _mdb.overview()})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/db/dates")
def db_dates_api():
    """某张表按日期聚合的行数"""
    table = (request.args.get("table") or "").strip()
    try:
        limit = int(request.args.get("limit", 120))
    except Exception:
        limit = 120
    try:
        return jsonify({"success": True, **_mdb.table_dates(table, limit)})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/db/cleanup", methods=["POST"])
def db_cleanup_api():
    """清理数据。scope: minute | table | before | all"""
    body = request.get_json(silent=True) or {}
    scope = (body.get("scope") or request.args.get("scope") or "").strip()
    table = (body.get("table") or request.args.get("table") or "").strip() or None
    before_date = (body.get("before_date") or request.args.get("before_date") or "").strip() or None
    keep_days = body.get("keep_days") or request.args.get("keep_days")
    if not scope:
        return jsonify({"success": False, "error": "missing param: scope"}), 400
    res = _mdb.cleanup(scope=scope, table=table, before_date=before_date, keep_days=keep_days)
    if res is None:
        return jsonify({"success": False, "error": "数据库不可用"}), 500
    if res.get("error"):
        return jsonify({"success": False, "error": res["error"]}), 400
    _mdb.log_fetch("db_cleanup", source="admin", status="ok", detail=str(res))
    return jsonify({"success": True, **res})


_JOB_FLAG = [False]
_JOB_NAME = [""]


@app.route("/api/db/run-job", methods=["POST"])
def db_run_job_api():
    """手动触发"每日抓取入库"任务（后台执行，立即返回）"""
    try:
        import scheduler as _sched
    except Exception as e:
        return jsonify({"success": False, "error": "scheduler unavailable: %s" % e}), 500
    if _JOB_FLAG[0]:
        return jsonify({"success": False, "error": "任务正在执行中，请稍候"}), 409
    body = request.get_json(silent=True) or {}
    target = (body.get("date") or request.args.get("date") or "").strip() or None
    job = (body.get("job") or request.args.get("job") or "close").strip() or "close"
    mode = (body.get("mode") or request.args.get("mode") or "").strip()
    if mode not in ("light", "full"):
        mode = "light" if job == "morning" else "full"
    _JOB_FLAG[0] = True
    _JOB_NAME[0] = "%s(%s)" % (job, mode)

    def _run():
        try:
            _sched.run_job(job, mode, target)
        except Exception as e:
            print("[run-job] %s" % e)
        finally:
            _JOB_FLAG[0] = False
            _JOB_NAME[0] = ""

    _threading.Thread(target=_run, daemon=True).start()
    return jsonify({"success": True, "started": True, "target": target or "auto",
                    "job": job, "mode": mode})


@app.route("/api/db/job-status")
def db_job_status_api():
    """定时任务状态：开关、时间、最近执行记录"""
    try:
        import scheduler as _sched
        running = _sched._THREAD is not None and _sched._THREAD.is_alive()
    except Exception:
        running = False
    info = _mdb.auto_fetch_info()
    info["thread_alive"] = running
    info["manual_running"] = _JOB_FLAG[0]
    info["running_job"] = _JOB_NAME[0]
    info["history"] = _mdb.recent_job_logs(8)
    info["pool"] = _mdb.pool_stats()
    return jsonify({"success": True, "job": info})


def _build_csv(headers, rows):
    """生成 CSV（带 UTF-8 BOM，Excel 打开中文不乱码）"""
    import csv as _csv
    sio = io.StringIO()
    sio.write("\ufeff")
    w = _csv.writer(sio)
    w.writerow(headers)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    return io.BytesIO(sio.getvalue().encode("utf-8"))


def _build_xlsx(sheets):
    """生成多 sheet 的 xlsx；缺少 openpyxl 时返回 None"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter
    except Exception:
        return None
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for label, headers, rows in sheets:
        name = (label or "Sheet")[:28]
        base = name
        i = 2
        while name in used:
            name = "%s%d" % (base, i)
            i += 1
        used.add(name)
        ws = wb.create_sheet(title=name)
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append(["" if v is None else v for v in r])
        ws.freeze_panes = "A2"
        for idx, h in enumerate(headers, 1):
            width = max(9, min(40, len(str(h)) * 2 + 6))
            ws.column_dimensions[get_column_letter(idx)].width = width
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


@app.route("/api/db/export-info")
def db_export_info_api():
    """某交易日各表可导出行数"""
    date = (request.args.get("date") or "").strip()
    if not date:
        return jsonify({"success": False, "error": "missing param: date"}), 400
    try:
        return jsonify({"success": True, "date": date,
                        "tables": _mdb.export_summary(date)})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/db/export")
def db_export_api():
    """导出某交易日数据。参数：date, format=csv|xlsx, table(可选)"""
    date = (request.args.get("date") or "").strip()
    fmt = (request.args.get("format") or "csv").strip().lower()
    table = (request.args.get("table") or "").strip()
    if not date:
        return jsonify({"success": False, "error": "missing param: date"}), 400
    try:
        if table:
            d = _mdb.export_rows(table, date)
            if d is None:
                return jsonify({"success": False, "error": "不支持导出的表: %s" % table}), 400
            if not d["rows"]:
                return jsonify({"success": False, "error": "%s 在 %s 没有数据" % (table, date)}), 404
            sheets = [(d["label"], d["headers"], d["rows"])]
            base = "%s_%s" % (date, d["label"])
        else:
            sheets = _mdb.export_all(date)
            base = "%s_全部数据" % date
        if not sheets:
            return jsonify({"success": False, "error": "%s 没有可导出的数据" % date}), 404

        if fmt == "xlsx":
            bio = _build_xlsx(sheets)
            if bio is None:
                return jsonify({"success": False, "error": "未安装 openpyxl，请改用 CSV"}), 500
            return send_file(
                bio, as_attachment=True, download_name=base + ".xlsx",
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        label, headers, rows = sheets[0]
        bio = _build_csv(headers, rows)
        return send_file(bio, as_attachment=True, download_name=base + ".csv",
                         mimetype="text/csv; charset=utf-8")
    except Exception as e:
        return jsonify({"success": False, "error": "导出失败: %s" % e}), 500


_BF_STATE = {"running": False, "progress": "", "result": None, "started_at": None, "finished_at": None}


def _load_backfill_module():
    """加载回补工具。

    开发时在 tools/ 下；打包后该文件被放在解包目录（因为它是运行时动态加载的，
    PyInstaller 追踪不到，必须显式随包带上）。
    """
    import importlib
    import importlib.util

    # 1) 常规 import（开发模式，tools 已在 sys.path）
    try:
        return importlib.import_module("backfill_hotlist")
    except Exception:
        pass

    # 2) 从解包目录 / 脚本目录直接按文件加载
    cands = []
    meipass = getattr(sys, "_MEIPASS", "")
    if meipass:
        cands.append(os.path.join(meipass, "backfill_hotlist.py"))
    here = os.path.dirname(os.path.abspath(__file__))
    cands.append(os.path.join(here, "backfill_hotlist.py"))
    cands.append(os.path.join(here, "tools", "backfill_hotlist.py"))
    for p in cands:
        if p and os.path.isfile(p):
            spec = importlib.util.spec_from_file_location("backfill_hotlist", p)
            mod = importlib.util.module_from_spec(spec)
            sys.modules["backfill_hotlist"] = mod
            spec.loader.exec_module(mod)
            return mod
    raise RuntimeError("找不到 backfill_hotlist.py")


@app.route("/api/themes/dates")
def themes_dates_api():
    """库里有题材数据的交易日（供题材页点选）"""
    try:
        limit = int(request.args.get("limit", 60))
    except Exception:
        limit = 60
    limit = max(1, min(limit, 200))
    try:
        return jsonify({"success": True, "dates": _mdb.theme_dates(limit)})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/hotlist/range")
def hotlist_range_api():
    """区间内每个交易日的人气榜数据情况（供"日期范围"浏览）"""
    start = (request.args.get("start") or "").strip()
    end = (request.args.get("end") or "").strip()
    if not start or not end:
        return jsonify({"success": False, "error": "missing param: start/end"}), 400
    try:
        rows = _mdb.hotlist_range_status(start, end)
        return jsonify({"success": True, "start": start, "end": end, "days": rows,
                        "summary": _mdb.hotlist_days_summary(start, end)})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/backfill/missing")
def backfill_missing_api():
    """列出最近 N 个交易日中缺人气榜数据的日期"""
    try:
        days = int(request.args.get("days", 20))
    except Exception:
        days = 20
    days = max(1, min(days, 120))
    try:
        bf = _load_backfill_module()
        st = bf.hotlist_stats(days)
        missing = st["missing"]
        # 区分：超出接口可回溯范围的（永远补不了）vs 还能补的
        lo = hi = None
        out_of_range = []
        try:
            lo, hi = bf.api_date_range()
        except Exception:
            pass
        if lo:
            out_of_range = [d for d in missing if d < lo]
            missing = [d for d in missing if d >= lo]
        return jsonify({"success": True, "days": days,
                        "dates": missing,
                        "missing": missing,
                        "out_of_range": out_of_range,
                        "api_from": lo, "api_to": hi,
                        "backfilled": st["backfilled"],
                        "live": st["live"]})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/backfill/run", methods=["POST"])
def backfill_run_api():
    """启动历史回补（后台执行，立即返回）"""
    if _BF_STATE["running"]:
        return jsonify({"success": False, "error": "回补任务正在执行中，请稍候"}), 409
    body = request.get_json(silent=True) or {}
    try:
        days = int(body.get("days") or 20)
    except Exception:
        days = 20
    days = max(1, min(days, 120))
    dates = [d for d in (body.get("dates") or []) if isinstance(d, str)]

    _BF_STATE.update(running=True, progress="准备中…", result=None,
                     started_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                     finished_at=None)

    def _run():
        try:
            bf = _load_backfill_module()
            targets = dates or bf._missing_trading_days(days)
            if not targets:
                _BF_STATE.update(running=False, progress="",
                                 result={"result": {}, "note": "所有交易日都已有数据，无需回补"})
                return
            # 接口只覆盖最近约 120 个自然日，过滤掉拿不到的日期
            lo, hi = bf.api_date_range()
            skipped = []
            if lo:
                skipped = [d for d in targets if d < lo]
                targets = [d for d in targets if d >= lo]
            if not targets:
                _BF_STATE.update(running=False, progress="",
                                 result={"result": {},
                                         "note": "这些日期都超出接口可回溯范围（最早 %s），无法回补" % lo})
                return
            _BF_STATE["progress"] = "将回补 %d 天，正在获取股票列表…%s" % (
                len(targets), ("（%d 天超出范围已跳过）" % len(skipped)) if skipped else "")

            def prog(stage, info):
                if stage == "universe":
                    _BF_STATE["progress"] = "全市场 %d 只，开始扫描…" % info
                elif stage == "scan":
                    try:
                        done, total = info
                        _BF_STATE["progress"] = "已扫描 %d/%d" % (done, total)
                    except Exception:
                        pass
            res = bf.backfill(targets, progress=prog)
            _BF_STATE.update(running=False, progress="", result=res)
        except Exception as e:
            _BF_STATE.update(running=False, progress="", result={"error": str(e)})
        finally:
            _BF_STATE["finished_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    _threading.Thread(target=_run, daemon=True).start()
    return jsonify({"success": True, "started": True, "days": days})


@app.route("/api/backfill/status")
def backfill_status_api():
    return jsonify({"success": True, "state": _BF_STATE})


@app.route("/api/backfill/clear", methods=["POST"])
def backfill_clear_api():
    """清除回补数据（不影响实时抓取的数据）"""
    body = request.get_json(silent=True) or {}
    d = (body.get("date") or request.args.get("date") or "").strip() or None
    n = _mdb.clear_backfill_hotlist(d)
    return jsonify({"success": True, "deleted": n})


@app.route("/api/db/coverage")
def db_coverage_api():
    """最近 N 个交易日的覆盖情况（哪些天有数据）"""
    try:
        n = int(request.args.get("n", 30))
    except Exception:
        n = 30
    n = max(5, min(n, 120))
    try:
        return jsonify({"success": True, "coverage": _mdb.calendar_coverage(n)})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/db/integrity")
def db_integrity_api():
    """数据完整性校验（默认校验最近交易日）"""
    target = (request.args.get("date") or "").strip() or None
    try:
        return jsonify({"success": True, "integrity": _mdb.integrity_check(target)})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/notify/status")
def notify_status_api():
    """通知通道配置概览"""
    try:
        import notifier
        return jsonify({"success": True, "notify": notifier.status()})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/notify/test", methods=["POST"])
def notify_test_api():
    """发送一条测试通知"""
    try:
        import notifier
    except Exception as e:
        return jsonify({"success": False, "error": "notifier 不可用: %s" % e}), 500
    chans = notifier.channels()
    if not chans:
        return jsonify({"success": False,
                        "error": "尚未配置任何通知通道（可在 exe 同目录创建 notify.json）"}), 400
    try:
        title = "【A股人气雷达】通知测试"
        text = ("这是一条测试消息。\n若你收到它，说明通知配置成功。\n"
                "发送时间：%s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        res = notifier.send(title, text, only=chans)
        return jsonify({"success": True, "results": res})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/db-stats")
def db_stats_api():
    """本地数据仓库概览（用于确认数据已落库）"""
    try:
        return jsonify({"success": True, "stats": _mdb.stats()})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


# ---- 启动每日定时抓取入库（AUTO_FETCH_ENABLED=0 可关闭）----
try:
    import scheduler as _scheduler
    _scheduler.start_if_enabled()
except Exception as _e:
    print("[scheduler] 启动失败: %s" % _e)


if __name__ == '__main__':
    print("=" * 50)
    print("A股人气雷达 服务已启动")
    print("访问地址: http://localhost:5000")
    print("=" * 50)
    app.run(host='0.0.0.0', port=5000, debug=True)







